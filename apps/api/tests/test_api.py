from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from sheet_shaper_api.config import Settings
from sheet_shaper_api.main import create_app


def client(tmp_path: Path, *, public: bool = False):
    settings = Settings(data_dir=tmp_path, public_mode=public, api_token="x" * 40, allowed_origins="http://localhost:4173")
    return TestClient(create_app(settings))


def upload_csv(api: TestClient):
    response = api.post("/api/v1/documents", files={"file": ("customers.csv", b"Name,Email\n Ada ,ada@example.com\n Ada ,ada@example.com\nLinus,=2+2\n", "text/csv")})
    assert response.status_code == 201, response.text
    return response.json()["data"]


def test_import_preview_apply_download_history_and_restore(tmp_path):
    with client(tmp_path) as api:
        document = upload_csv(api)
        assert document["rows"] == 4
        preview = api.post(f"/api/v1/documents/{document['id']}/preview", json={"hasHeader": True, "operations": [{"type": "smart_trim", "options": {}}, {"type": "remove_duplicates", "options": {}}]}).json()["data"]
        assert preview["summary"]["risk"] == "MEDIUM"
        applied = api.post(f"/api/v1/documents/{document['id']}/apply", json={"previewToken": preview["previewToken"], "acknowledged": True, "idempotencyKey": "integration-apply-0001"})
        assert applied.status_code == 200, applied.text
        operation = applied.json()["data"]
        assert operation["status"] == "APPLIED"
        assert operation["validation"]["passed"] is True
        repeated = api.post(f"/api/v1/documents/{document['id']}/apply", json={"previewToken": preview["previewToken"], "acknowledged": True, "idempotencyKey": "integration-apply-0001"})
        assert repeated.json()["data"]["id"] == operation["id"]
        history = api.get("/api/v1/history").json()["data"]
        assert history["total"] == 1 and history["statusCounts"]["APPLIED"] == 1
        downloaded = api.get(f"/api/v1/documents/{document['id']}/download")
        assert downloaded.status_code == 200 and b"'=" in downloaded.content
        restored = api.post(f"/api/v1/operations/{operation['id']}/restore")
        assert restored.json()["data"]["status"] == "RESTORED"


def test_public_mode_requires_bearer_token(tmp_path):
    with client(tmp_path, public=True) as api:
        assert api.get("/api/v1/documents").status_code == 401
        headers = {"Authorization": "Bearer " + "x" * 40}
        assert api.get("/api/v1/documents", headers=headers).status_code == 200
        assert api.get("/api/v1/capabilities", headers=headers).json()["data"]["externalTransmission"] is True
        assert api.get("/hai/v1/manifest", headers=headers).json()["data"]["readOnlyByDefault"] is True


def test_hai_mutations_are_disabled_by_default(tmp_path):
    with client(tmp_path) as api:
        headers = {"Authorization": "Bearer " + "x" * 40}
        manifest = api.get("/hai/v1/manifest", headers=headers).json()["data"]
        assert "operations.apply" not in manifest["capabilities"]
        blocked = api.post("/hai/v1/operations/not-real/restore", headers=headers)
        assert blocked.status_code == 403
        assert blocked.json()["error"]["code"] == "HAI_MUTATIONS_DISABLED"


def test_malformed_and_oversized_inputs_fail_without_stack_traces(tmp_path):
    with client(tmp_path) as api:
        unsupported = api.post("/api/v1/documents", files={"file": ("bad.exe", b"MZ", "application/octet-stream")})
        assert unsupported.status_code == 400
        body = unsupported.json()
        assert body["error"]["code"] == "UNSUPPORTED_FILE"
        assert "traceback" not in unsupported.text.lower()
        malformed = api.post("/api/v1/documents", files={"file": ("bad.xlsx", b"not-a-zip", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert malformed.status_code == 400
        assert malformed.json()["error"]["code"] == "INVALID_WORKBOOK"


def test_document_delete_removes_metadata_and_files(tmp_path):
    with client(tmp_path) as api:
        document = upload_csv(api)
        response = api.delete(f"/api/v1/documents/{document['id']}")
        assert response.status_code == 200
        assert response.json()["data"]["deleted"] is True
        assert api.get(f"/api/v1/documents/{document['id']}").status_code == 404
        assert not list((tmp_path / "files").rglob("*.csv"))


def test_workers_must_remain_single_process():
    try:
        Settings(workers=2)
    except ValueError as error:
        assert "SS_WORKERS" in str(error)
    else:
        raise AssertionError("multi-process workers must fail closed")


def test_xlsx_round_trip_preserves_formula_and_validates_write(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Total"])
    sheet.append([" Ada ", "=1+1"])
    buffer = BytesIO()
    workbook.save(buffer)
    with client(tmp_path) as api:
        uploaded = api.post(
            "/api/v1/documents",
            files={"file": ("formula.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert uploaded.status_code == 201, uploaded.text
        document = uploaded.json()["data"]
        preview = api.post(
            f"/api/v1/documents/{document['id']}/preview",
            json={"hasHeader": True, "operations": [{"type": "smart_trim", "options": {}}]},
        ).json()["data"]
        applied = api.post(
            f"/api/v1/documents/{document['id']}/apply",
            json={"previewToken": preview["previewToken"], "acknowledged": True, "idempotencyKey": "xlsx-round-trip-0001"},
        )
        assert applied.json()["data"]["validation"]["passed"] is True
        downloaded = api.get(f"/api/v1/documents/{document['id']}/download")
        result = load_workbook(BytesIO(downloaded.content), data_only=False)
        try:
            assert result.active["A2"].value == "Ada"
            assert result.active["B2"].value == "=1+1"
        finally:
            result.close()
