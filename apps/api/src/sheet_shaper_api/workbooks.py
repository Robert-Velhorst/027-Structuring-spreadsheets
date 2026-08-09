from __future__ import annotations

import csv
import hashlib
import shutil
from zipfile import BadZipFile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils.cell import get_column_letter, range_boundaries

from .engine import EngineError, Table


CSV_TYPES = {"text/csv", "application/csv", "text/plain"}
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SUPPORTED_SUFFIXES = {".csv": "text/csv", ".xlsx": XLSX_TYPE}


@dataclass(frozen=True)
class WorkbookInfo:
    rows: int
    columns: int
    sheet_name: str
    media_type: str
    delimiter: str | None = None


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def detect_media_type(filename: str, declared: str | None) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise EngineError("UNSUPPORTED_FILE", "Only .csv and .xlsx files are supported.")
    expected = SUPPORTED_SUFFIXES[suffix]
    if expected == XLSX_TYPE and declared and declared not in {XLSX_TYPE, "application/octet-stream"}:
        raise EngineError("MEDIA_TYPE_MISMATCH", "The uploaded content type does not match an .xlsx file.")
    return expected


def safe_download_name(name: str) -> str:
    cleaned = "".join(character for character in name if character.isalnum() or character in " ._-").strip(" .")
    return (cleaned or "sheet-shaper-output")[:120]


class WorkbookAdapter:
    def __init__(self, path: Path, media_type: str):
        self.path = path
        self.media_type = media_type
        self._delimiter: str | None = None

    def inspect(self) -> WorkbookInfo:
        table, sheet_name, delimiter, _ = self.read("")
        return WorkbookInfo(table.height, table.width, sheet_name, self.media_type, delimiter)

    def read(self, target_range: str) -> tuple[Table, str, str | None, str]:
        if self.media_type in CSV_TYPES:
            return self._read_csv(target_range)
        return self._read_xlsx(target_range)

    def write(self, destination: Path, target_range: str, result: Table) -> WorkbookInfo:
        destination.parent.mkdir(parents=True, exist_ok=True)
        if self.media_type in CSV_TYPES:
            return self._write_csv(destination, target_range, result)
        return self._write_xlsx(destination, target_range, result)

    def prepare_for_write(self, table: Table) -> Table:
        rectangular = table.rectangular()
        return Table([
            [_safe_literal(value, (r, c) in rectangular.formulas) for c, value in enumerate(row)]
            for r, row in enumerate(rectangular.rows)
        ], rectangular.formulas)

    def _read_csv(self, target_range: str) -> tuple[Table, str, str, str]:
        try:
            with self.path.open("r", encoding="utf-8-sig", newline="") as stream:
                sample = stream.read(8192); stream.seek(0)
                try: delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
                except csv.Error: delimiter = ","
                rows = [list(row) for row in csv.reader(stream, delimiter=delimiter)]
        except UnicodeDecodeError as error:
            raise EngineError("INVALID_WORKBOOK", "The CSV must be valid UTF-8 text.") from error
        table = Table(rows).rectangular()
        normalized_range = _normalize_range(target_range, table.height, table.width)
        if target_range:
            min_col, min_row, max_col, max_row = range_boundaries(normalized_range)
            table = Table([row[min_col - 1:max_col] for row in table.rows[min_row - 1:max_row]])
        return table, "CSV", delimiter, normalized_range

    def _read_xlsx(self, target_range: str) -> tuple[Table, str, None, str]:
        try:
            workbook = load_workbook(self.path, read_only=True, data_only=False)
        except (BadZipFile, InvalidFileException, KeyError, OSError) as error:
            raise EngineError("INVALID_WORKBOOK", "The XLSX file is malformed or unreadable.") from error
        try:
            sheet = workbook.active
            max_row, max_column = max(sheet.max_row, 1), max(sheet.max_column, 1)
            normalized_range = _normalize_range(target_range, max_row, max_column)
            min_col, min_row, max_col, max_row = range_boundaries(normalized_range)
            rows: list[list[Any]] = []
            formulas: set[tuple[int, int]] = set()
            for output_row, values in enumerate(sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=False)):
                row: list[Any] = []
                for output_column, cell in enumerate(values):
                    row.append(cell.value if cell.value is not None else "")
                    if cell.data_type == "f": formulas.add((output_row, output_column))
                rows.append(row)
            return Table(rows, frozenset(formulas)).rectangular(), sheet.title, None, normalized_range
        finally:
            workbook.close()

    def _write_csv(self, destination: Path, target_range: str, result: Table) -> WorkbookInfo:
        full, _, delimiter, _ = self._read_csv("")
        normalized = _normalize_range(target_range, full.height, full.width) if target_range else ""
        updated = _patch_table(full, normalized, result)
        with destination.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, delimiter=delimiter)
            for r, row in enumerate(updated.rows):
                writer.writerow([value for value in row])
        return WorkbookInfo(updated.height, updated.width, "CSV", self.media_type, delimiter)

    def _write_xlsx(self, destination: Path, target_range: str, result: Table) -> WorkbookInfo:
        workbook = load_workbook(self.path, read_only=False, data_only=False)
        try:
            sheet = workbook.active
            normalized = _normalize_range(target_range, max(sheet.max_row, 1), max(sheet.max_column, 1))
            min_col, min_row, max_col, max_row = range_boundaries(normalized)
            clear_rows = max(max_row - min_row + 1, result.height)
            clear_columns = max(max_col - min_col + 1, result.width)
            for row in sheet.iter_rows(min_row=min_row, max_row=min_row + clear_rows - 1, min_col=min_col, max_col=min_col + clear_columns - 1):
                for cell in row: cell.value = None
            for r, row in enumerate(result.rows):
                for c, value in enumerate(row):
                    sheet.cell(min_row + r, min_col + c).value = value
            workbook.save(destination)
            return WorkbookInfo(sheet.max_row, sheet.max_column, sheet.title, self.media_type)
        finally:
            workbook.close()


def _normalize_range(target_range: str, rows: int, columns: int) -> str:
    if not target_range:
        return f"A1:{get_column_letter(max(columns, 1))}{max(rows, 1)}"
    try:
        min_col, min_row, max_col, max_row = range_boundaries(target_range)
    except ValueError as error:
        raise EngineError("INVALID_RANGE", "Use an A1 range such as A1:H100.") from error
    if min_row < 1 or min_col < 1 or max_row > rows or max_col > columns:
        raise EngineError("RANGE_OUT_OF_BOUNDS", "The target range must stay inside the used spreadsheet area.")
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def _patch_table(full: Table, target_range: str, result: Table) -> Table:
    if not target_range:
        return result.rectangular()
    min_col, min_row, max_col, max_row = range_boundaries(target_range)
    rows = [list(row) for row in full.rectangular().rows]
    required_rows = min_row - 1 + result.height
    required_columns = min_col - 1 + result.width
    width = max(full.width, required_columns)
    while len(rows) < required_rows: rows.append([""] * width)
    rows = [row + [""] * (width - len(row)) for row in rows]
    clear_rows = max(max_row - min_row + 1, result.height); clear_columns = max(max_col - min_col + 1, result.width)
    for r in range(clear_rows):
        for c in range(clear_columns): rows[min_row - 1 + r][min_col - 1 + c] = ""
    for r, row in enumerate(result.rows):
        for c, value in enumerate(row): rows[min_row - 1 + r][min_col - 1 + c] = value
    return Table(rows)


def _safe_literal(value: Any, formula: bool) -> Any:
    if formula: return value
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")): return "'" + value
    return value


def copy_file(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with source.open("rb") as input_stream, destination.open("xb") as output_stream:
        shutil.copyfileobj(input_stream, output_stream, 1024 * 1024)
