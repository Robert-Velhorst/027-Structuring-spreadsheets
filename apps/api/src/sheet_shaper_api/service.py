from __future__ import annotations

import asyncio
import json
import os
import re
import threading
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl.utils.cell import get_column_letter, range_boundaries

from .config import Settings
from .db import Database
from .engine import OPERATION_RISK, EngineError, Table, apply_pipeline, difference, digest, normalize_request, quality
from .workbooks import WorkbookAdapter, copy_file, detect_media_type, safe_download_name, sha256_file


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def identifier() -> str:
    return str(uuid.uuid4())


def json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=str)


class ServiceError(RuntimeError):
    def __init__(self, code: str, message: str, status_code: int = 400, retry_safe: bool = False):
        super().__init__(message)
        self.code = code
        self.status_code = status_code
        self.retry_safe = retry_safe


class SheetShaperService:
    def __init__(self, settings: Settings, database: Database):
        self.settings = settings
        self.database = database
        self._locks_guard = threading.Lock()
        self._document_locks: dict[str, threading.Lock] = {}

    def initialize(self) -> None:
        self.settings.files_dir.mkdir(parents=True, exist_ok=True)
        self.database.initialize()

    def prune_expired_metadata(self) -> None:
        timestamp = now()
        with self.database.transaction() as connection:
            connection.execute("DELETE FROM previews WHERE expires_at < ?", (timestamp,))
            connection.execute(
                """DELETE FROM audit_events WHERE owner=? AND id NOT IN (
                    SELECT id FROM audit_events WHERE owner=? ORDER BY id DESC LIMIT ?
                )""",
                ("local-owner", "local-owner", self.settings.max_audit_events),
            )

    def _lock(self, document_id: str) -> threading.Lock:
        with self._locks_guard:
            return self._document_locks.setdefault(document_id, threading.Lock())

    def file_path(self, file_id: str, suffix: str) -> Path:
        if suffix not in {".csv", ".xlsx"}: raise ServiceError("UNSUPPORTED_FILE", "Unsupported file extension.")
        return self.settings.files_dir / file_id[:2] / f"{file_id}{suffix}"

    async def import_upload(self, upload: Any, owner: str) -> dict[str, Any]:
        filename = safe_download_name(upload.filename or "spreadsheet")
        media_type = detect_media_type(filename, upload.content_type)
        suffix = Path(filename).suffix.lower(); file_id = identifier(); path = self.file_path(file_id, suffix); path.parent.mkdir(parents=True, exist_ok=True)
        limit = self.settings.max_upload_mb * 1024 * 1024; size = 0
        try:
            with path.open("xb") as stream:
                while chunk := await upload.read(1024 * 1024):
                    size += len(chunk)
                    if size > limit: raise ServiceError("UPLOAD_TOO_LARGE", f"Uploads are limited to {self.settings.max_upload_mb} MB.", 413)
                    stream.write(chunk)
            if size == 0: raise ServiceError("EMPTY_UPLOAD", "The uploaded file is empty.")
            adapter = WorkbookAdapter(path, media_type); info = await asyncio.to_thread(adapter.inspect)
            if info.rows * info.columns > self.settings.max_cells: raise ServiceError("RANGE_TOO_LARGE", f"The workbook contains more than {self.settings.max_cells:,} used cells.", 413)
            document_id, version_id, timestamp = identifier(), identifier(), now()
            source_hash = await asyncio.to_thread(sha256_file, path)
            with self.database.transaction() as connection:
                connection.execute("INSERT INTO documents(id,owner,name,media_type,current_version_id,created_at,updated_at) VALUES (?,?,?,?,?,?,?)", (document_id, owner, filename, media_type, version_id, timestamp, timestamp))
                connection.execute("INSERT INTO versions(id,document_id,kind,file_path,sha256,rows,columns,created_at) VALUES (?,?,?,?,?,?,?,?)", (version_id, document_id, "SOURCE", str(path), source_hash, info.rows, info.columns, timestamp))
                self._audit(connection, owner, "document.import", document_id, None, "success", {"mediaType": media_type, "bytes": size, "rows": info.rows, "columns": info.columns})
            return await asyncio.to_thread(self.get_document, document_id, owner, True)
        except Exception:
            path.unlink(missing_ok=True)
            raise
        finally:
            await upload.close()

    def list_documents(self, owner: str) -> list[dict[str, Any]]:
        with self.database.connection() as connection:
            rows = connection.execute("""SELECT d.*,v.rows,v.columns,v.sha256 FROM documents d
                JOIN versions v ON v.id=d.current_version_id WHERE d.owner=? ORDER BY d.updated_at DESC LIMIT 100""", (owner,)).fetchall()
        return [self._document_public(row) for row in rows]

    def get_document(self, document_id: str, owner: str, include_window: bool = False, row: int = 1, column: int = 1, rows: int = 100, columns: int = 30) -> dict[str, Any]:
        document, version = self._document_and_version(document_id, owner)
        result = self._document_public({**dict(document), **{"rows": version["rows"], "columns": version["columns"], "sha256": version["sha256"]}})
        if include_window:
            adapter = WorkbookAdapter(Path(version["file_path"]), document["media_type"]); table, sheet, _, target = adapter.read("")
            start_row, start_column = max(row - 1, 0), max(column - 1, 0)
            window = [values[start_column:start_column + columns] for values in table.rows[start_row:start_row + rows]]
            result["window"] = {"row": row, "column": column, "rows": window, "sheet": sheet, "range": target, "totalRows": table.height, "totalColumns": table.width, "truncated": table.height > start_row + rows or table.width > start_column + columns}
            result["quality"] = quality(table)
        return result

    def preview(self, document_id: str, request: dict[str, Any], owner: str) -> dict[str, Any]:
        document, version = self._document_and_version(document_id, owner); normalized = normalize_request(request)
        adapter = WorkbookAdapter(Path(version["file_path"]), document["media_type"]); before, sheet, _, target = adapter.read(normalized["range"])
        if before.height * before.width > self.settings.max_cells: raise ServiceError("RANGE_TOO_LARGE", f"Select no more than {self.settings.max_cells:,} cells.", 413)
        after = adapter.prepare_for_write(apply_pipeline(before, normalized))
        if not after.height or not after.width: raise ServiceError("EMPTY_RESULT", "The operation sequence would leave an empty target.")
        summary = difference(before, after); summary["risk"] = self._request_risk(normalized); summary["target"] = f"{sheet}!{target}"; summary["beforeQuality"] = quality(before, normalized["hasHeader"]); summary["afterQuality"] = quality(after, normalized["hasHeader"])
        preview_id, timestamp = identifier(), now(); expires = (datetime.now(timezone.utc) + timedelta(seconds=self.settings.preview_ttl_seconds)).isoformat()
        normalized["range"] = target
        with self.database.transaction() as connection:
            connection.execute("INSERT INTO previews(id,document_id,version_id,owner,request_json,input_digest,summary_json,expires_at,created_at) VALUES (?,?,?,?,?,?,?,?,?)", (preview_id, document_id, version["id"], owner, json_text(normalized), digest(before), json_text(summary), expires, timestamp))
            connection.execute("DELETE FROM previews WHERE expires_at < ? AND consumed_at IS NOT NULL", ((datetime.now(timezone.utc) - timedelta(days=1)).isoformat(),))
            self._audit(connection, owner, "operation.preview", document_id, None, "success", {"previewId": preview_id, "changedCells": summary["changedCells"], "risk": summary["risk"]})
        return {"previewToken": preview_id, "expiresAt": expires, "target": summary["target"], "summary": summary, "examples": summary["examples"], "warnings": ["No changes are applied until Apply once is confirmed.", "A durable backup is created and hashed before every write."]}

    def apply(self, document_id: str, command: dict[str, Any], owner: str) -> dict[str, Any]:
        if command.get("acknowledged") is not True: raise ServiceError("ACKNOWLEDGEMENT_REQUIRED", "Review and acknowledge the preview before applying.")
        preview_id = str(command.get("previewToken") or ""); idem = str(command.get("idempotencyKey") or "")
        with self.database.connection() as connection:
            prior = connection.execute("SELECT * FROM operations WHERE owner=? AND idempotency_key=?", (owner, idem)).fetchone()
            if prior: return self._operation_public(prior, connection)
        lock = self._lock(document_id)
        if not lock.acquire(timeout=5): raise ServiceError("BUSY", "Another write is active for this document.", 409, True)
        operation_id = identifier(); backup_path: Path | None = None
        try:
            document, current = self._document_and_version(document_id, owner)
            with self.database.connection() as connection:
                preview = connection.execute("SELECT * FROM previews WHERE id=? AND document_id=? AND owner=?", (preview_id, document_id, owner)).fetchone()
            if not preview: raise ServiceError("PREVIEW_NOT_FOUND", "Create a fresh preview.", 404)
            if preview["consumed_at"]: raise ServiceError("PREVIEW_USED", "This preview was already used. Create a fresh preview.", 409)
            if datetime.fromisoformat(preview["expires_at"]) <= datetime.now(timezone.utc): raise ServiceError("PREVIEW_EXPIRED", "The preview expired. Create a fresh preview.", 409, True)
            if preview["version_id"] != current["id"]: raise ServiceError("STALE_PREVIEW", "The document changed after preview. No changes were made.", 409, True)
            request = json.loads(preview["request_json"]); adapter = WorkbookAdapter(Path(current["file_path"]), document["media_type"]); before, _, _, target = adapter.read(request["range"])
            if digest(before) != preview["input_digest"]: raise ServiceError("STALE_PREVIEW", "The target changed after preview. No changes were made.", 409, True)
            after = adapter.prepare_for_write(apply_pipeline(before, request)); suffix = Path(current["file_path"]).suffix
            backup_id, result_id = identifier(), identifier(); backup_path = self.file_path(backup_id, suffix); result_path = self.file_path(result_id, suffix)
            copy_file(Path(current["file_path"]), backup_path); backup_hash = sha256_file(backup_path)
            if backup_hash != current["sha256"]: raise ServiceError("BACKUP_VALIDATION_FAILED", "The backup hash did not match the source. No changes were made.", 500)
            timestamp = now(); summary = json.loads(preview["summary_json"])
            with self.database.transaction() as connection:
                connection.execute("INSERT INTO versions(id,document_id,kind,file_path,sha256,rows,columns,created_at) VALUES (?,?,?,?,?,?,?,?)", (backup_id, document_id, "BACKUP", str(backup_path), backup_hash, current["rows"], current["columns"], timestamp))
                connection.execute("INSERT INTO operations(id,document_id,preview_id,owner,before_version_id,backup_version_id,status,summary_json,idempotency_key,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)", (operation_id, document_id, preview_id, owner, current["id"], backup_id, "APPLYING", json_text(summary), idem, timestamp, timestamp))
            try:
                info = adapter.write(result_path, target, after); result_hash = sha256_file(result_path)
                validation_adapter = WorkbookAdapter(result_path, document["media_type"]); actual, _, _, _ = validation_adapter.read(_output_range(target, after))
                expected_digest, actual_digest = digest(after), digest(actual)
                validation = {"passed": expected_digest == actual_digest, "expectedDigest": expected_digest, "actualDigest": actual_digest, "checkedAt": now(), "quality": quality(actual, request["hasHeader"])}
                if not validation["passed"]: raise ServiceError("VALIDATION_FAILED", "The written target did not match the preview. Restore is available.", 500)
                completed = now()
                with self.database.transaction() as connection:
                    connection.execute("INSERT INTO versions(id,document_id,kind,file_path,sha256,rows,columns,created_at) VALUES (?,?,?,?,?,?,?,?)", (result_id, document_id, "RESULT", str(result_path), result_hash, info.rows, info.columns, completed))
                    connection.execute("UPDATE operations SET after_version_id=?,status='APPLIED',validation_json=?,updated_at=? WHERE id=?", (result_id, json_text(validation), completed, operation_id))
                    connection.execute("UPDATE previews SET consumed_at=? WHERE id=?", (completed, preview_id))
                    connection.execute("UPDATE documents SET current_version_id=?,updated_at=? WHERE id=?", (result_id, completed, document_id))
                    self._audit(connection, owner, "operation.apply", document_id, operation_id, "success", {"changedCells": summary["changedCells"], "validationPassed": True})
                return self.get_operation(operation_id, owner)
            except Exception as error:
                result_path.unlink(missing_ok=True)
                with self.database.transaction() as connection:
                    connection.execute("UPDATE operations SET status='FAILED',error_code=?,updated_at=? WHERE id=?", (getattr(error, "code", "WRITE_FAILED"), now(), operation_id))
                    self._audit(connection, owner, "operation.apply", document_id, operation_id, "failure", {"errorCode": getattr(error, "code", "WRITE_FAILED")})
                raise
        finally:
            lock.release()

    def restore(self, operation_id: str, owner: str) -> dict[str, Any]:
        with self.database.connection() as connection:
            operation = connection.execute("SELECT * FROM operations WHERE id=? AND owner=?", (operation_id, owner)).fetchone()
            if not operation: raise ServiceError("NOT_FOUND", "Operation not found.", 404)
            document = connection.execute("SELECT * FROM documents WHERE id=? AND owner=?", (operation["document_id"], owner)).fetchone()
            backup = connection.execute("SELECT * FROM versions WHERE id=?", (operation["backup_version_id"],)).fetchone()
        if operation["status"] not in {"APPLIED", "FAILED"}: raise ServiceError("INVALID_STATE", "Only applied or failed operations can be restored.", 409)
        expected_current = operation["after_version_id"] if operation["status"] == "APPLIED" else operation["before_version_id"]
        if document["current_version_id"] != expected_current: raise ServiceError("STALE_RESTORE", "A newer document version exists. Restore that version first or download the backup.", 409)
        lock = self._lock(document["id"])
        if not lock.acquire(timeout=5): raise ServiceError("BUSY", "Another write is active for this document.", 409, True)
        try:
            suffix = Path(backup["file_path"]).suffix; restored_id = identifier(); restored_path = self.file_path(restored_id, suffix); copy_file(Path(backup["file_path"]), restored_path)
            restored_hash = sha256_file(restored_path)
            if restored_hash != backup["sha256"]: restored_path.unlink(missing_ok=True); raise ServiceError("RESTORE_VALIDATION_FAILED", "Restore hash validation failed. The backup remains available.", 500)
            timestamp = now()
            with self.database.transaction() as connection:
                connection.execute("INSERT INTO versions(id,document_id,kind,file_path,sha256,rows,columns,created_at) VALUES (?,?,?,?,?,?,?,?)", (restored_id, document["id"], "RESTORED", str(restored_path), restored_hash, backup["rows"], backup["columns"], timestamp))
                connection.execute("UPDATE documents SET current_version_id=?,updated_at=? WHERE id=?", (restored_id, timestamp, document["id"]))
                connection.execute("UPDATE operations SET status='RESTORED',updated_at=? WHERE id=?", (timestamp, operation_id))
                self._audit(connection, owner, "operation.restore", document["id"], operation_id, "success", {"restoredVersion": restored_id})
            return self.get_operation(operation_id, owner)
        finally:
            lock.release()

    def history(self, owner: str, page: int = 1, page_size: int = 20, status: str = "", query: str = "") -> dict[str, Any]:
        page, page_size = max(page, 1), min(max(page_size, 1), 50); filters = ["o.owner=?"]; values: list[Any] = [owner]
        if status: filters.append("o.status=?"); values.append(status.upper())
        if query: filters.append("(d.name LIKE ? OR o.summary_json LIKE ?)"); values.extend([f"%{query}%", f"%{query}%"])
        where = " AND ".join(filters)
        with self.database.connection() as connection:
            total = connection.execute(f"SELECT COUNT(*) FROM operations o JOIN documents d ON d.id=o.document_id WHERE {where}", values).fetchone()[0]
            records = connection.execute(f"SELECT o.*,d.name,d.current_version_id FROM operations o JOIN documents d ON d.id=o.document_id WHERE {where} ORDER BY o.created_at DESC LIMIT ? OFFSET ?", (*values, page_size, (page - 1) * page_size)).fetchall()
            counts = {row["status"]: row["count"] for row in connection.execute("SELECT status,COUNT(*) count FROM operations WHERE owner=? GROUP BY status", (owner,)).fetchall()}
            items = [self._operation_public(record, connection) for record in records]
        pages = max(1, (total + page_size - 1) // page_size)
        return {"items": items, "page": min(page, pages), "pageSize": page_size, "total": total, "pages": pages, "statusCounts": {key: counts.get(key, 0) for key in ["APPLIED", "FAILED", "RESTORED", "APPLYING"]}}

    def get_operation(self, operation_id: str, owner: str) -> dict[str, Any]:
        with self.database.connection() as connection:
            record = connection.execute("SELECT o.*,d.name,d.current_version_id FROM operations o JOIN documents d ON d.id=o.document_id WHERE o.id=? AND o.owner=?", (operation_id, owner)).fetchone()
            if not record: raise ServiceError("NOT_FOUND", "Operation not found.", 404)
            return self._operation_public(record, connection)

    def download(self, document_id: str, owner: str) -> tuple[Path, str, str]:
        document, version = self._document_and_version(document_id, owner)
        stem, suffix = Path(document["name"]).stem, Path(document["name"]).suffix
        return Path(version["file_path"]), safe_download_name(f"{stem}-shaped{suffix}"), document["media_type"]

    def delete_document(self, document_id: str, owner: str) -> None:
        lock = self._lock(document_id)
        if not lock.acquire(timeout=5):
            raise ServiceError("BUSY", "Another write is active for this document.", 409, True)
        try:
            with self.database.transaction() as connection:
                document = connection.execute("SELECT id FROM documents WHERE id=? AND owner=?", (document_id, owner)).fetchone()
                if not document:
                    raise ServiceError("NOT_FOUND", "Document not found.", 404)
                paths = [Path(row["file_path"]) for row in connection.execute("SELECT file_path FROM versions WHERE document_id=?", (document_id,)).fetchall()]
                self._audit(connection, owner, "document.delete", document_id, None, "success", {"versions": len(paths)})
                connection.execute("DELETE FROM documents WHERE id=? AND owner=?", (document_id, owner))
            for path in paths:
                path.unlink(missing_ok=True)
                try:
                    path.parent.rmdir()
                except OSError:
                    pass
        finally:
            lock.release()
            with self._locks_guard:
                self._document_locks.pop(document_id, None)

    def _document_and_version(self, document_id: str, owner: str) -> tuple[Any, Any]:
        with self.database.connection() as connection:
            document = connection.execute("SELECT * FROM documents WHERE id=? AND owner=?", (document_id, owner)).fetchone()
            if not document: raise ServiceError("NOT_FOUND", "Document not found.", 404)
            version = connection.execute("SELECT * FROM versions WHERE id=? AND document_id=?", (document["current_version_id"], document_id)).fetchone()
            if not version: raise ServiceError("VERSION_MISSING", "The current document version is unavailable.", 500)
            return document, version

    @staticmethod
    def _document_public(row: Any) -> dict[str, Any]:
        return {"id": row["id"], "name": row["name"], "mediaType": row["media_type"], "rows": row["rows"], "columns": row["columns"], "versionHash": row["sha256"][:16], "createdAt": row["created_at"], "updatedAt": row["updated_at"]}

    def _operation_public(self, row: Any, connection: Any) -> dict[str, Any]:
        summary = json.loads(row["summary_json"]); validation = json.loads(row["validation_json"]) if row["validation_json"] else None
        current = row["current_version_id"] if "current_version_id" in row.keys() else connection.execute("SELECT current_version_id FROM documents WHERE id=?", (row["document_id"],)).fetchone()[0]
        expected = row["after_version_id"] if row["status"] == "APPLIED" else row["before_version_id"]
        return {"id": row["id"], "documentId": row["document_id"], "documentName": row["name"] if "name" in row.keys() else "", "status": row["status"], "summary": summary, "validation": validation, "errorCode": row["error_code"], "canRestore": row["status"] in {"APPLIED", "FAILED"} and current == expected, "createdAt": row["created_at"], "updatedAt": row["updated_at"]}

    @staticmethod
    def _request_risk(request: dict[str, Any]) -> str:
        weights = {"LOW": 1, "MEDIUM": 2, "HIGH": 3}; return max((OPERATION_RISK[item["type"]] for item in request["operations"]), key=lambda value: weights[value])

    @staticmethod
    def _audit(connection: Any, owner: str, event: str, document_id: str | None, operation_id: str | None, outcome: str, detail: dict[str, Any]) -> None:
        safe_detail = {key: value for key, value in detail.items() if key not in {"rows", "cells", "content", "token"}}
        connection.execute("INSERT INTO audit_events(owner,event,document_id,operation_id,outcome,detail_json,created_at) VALUES (?,?,?,?,?,?,?)", (owner, event, document_id, operation_id, outcome, json_text(safe_detail), now()))


def _output_range(target: str, table: Table) -> str:
    min_col, min_row, _, _ = range_boundaries(target)
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(min_col + table.width - 1)}{min_row + table.height - 1}"
